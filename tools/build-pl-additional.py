#!/usr/bin/env python3
"""
Generate js/content-pl-additional.js — the long-form Polish section added below
each instrument page's existing content — from the approved workbook.

Source of truth: "Instrument_Pages_PL_Content_EURUSD_to_EWJ.xlsx"
  * Instrument Page Content PL -> one row per content block, in page order
      A No. · B Instrument · C Target Page URL · D Page Title (H1)
      E Content Type · F Content (PL) · G Link(s) in Content · H Status
  * Instrument Index           -> expected "Content Blocks" count per page

Column F is copied verbatim. Column C picks the page (its path is matched to a
CATALOGUE slug in js/data.js, ignoring hyphens: /xau-aud -> xauaud).

Column G comes in two shapes:
  * "anchor -> url; anchor -> url"  — anchor text given by the workbook
  * "url"                           — bare URL; its anchor is read from
                                      tools/link-anchors.json (ticker -> url -> text)
Every anchor must occur verbatim in that row's Column F, or the build fails.

The build also fails if a page's block count differs from the Instrument Index.

Usage:  python3 tools/build-pl-additional.py [workbook.xlsx] [--only EURUSD,GBPJPY]
"""

import argparse
import json
import os
import re
from urllib.parse import urlparse

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.expanduser("~/Downloads/Instrument_Pages_PL_Content_EURUSD_to_EWJ.xlsx")
OUT = os.path.join(HERE, "..", "js", "content-pl-additional.js")
DATA_JS = os.path.join(HERE, "..", "js", "data.js")
ANCHORS = os.path.join(HERE, "link-anchors.json")

# Column E -> block type rendered by js/additional-content.js
TYPES = {"H1": "h1", "H2": "h2", "H3": "h3", "Paragraph": "p", "Bullet": "li"}


def clean(v):
    return "" if v is None else str(v).strip()


def slug(s):
    return re.sub(r"(^-|-$)", "", re.sub(r"[^a-z0-9]+", "-", str(s).lower()))


def site_ids():
    """slug(symbol) for every CATALOGUE symbol in js/data.js."""
    src = open(DATA_JS, encoding="utf-8").read()
    block = re.search(r"var CATALOGUE = \{(.*?)\};", src, re.S).group(1)
    syms = re.findall(r'"([^"]+)"', re.sub(r'"[^"]+":\s*\[', "[", block))
    return {slug(s) for s in syms}


def page_id(url, ids):
    path = urlparse(url).path.strip("/").lower()
    if path in ids:
        return path
    flat = {i.replace("-", ""): i for i in ids}
    return flat.get(path.replace("-", ""))


def parse_links(cell, ticker, text, anchors):
    """Column G -> [{text, href}], each anchor verified against Column F."""
    out = []
    for part in [p.strip() for p in clean(cell).split(";") if p.strip()]:
        if "->" in part:
            anchor, href = [s.strip() for s in part.split("->", 1)]
        else:
            href = part
            anchor = anchors.get(ticker, {}).get(href)
            if not anchor:
                raise SystemExit("%s: no anchor for bare URL %s — add it to tools/link-anchors.json" % (ticker, href))
        if anchor not in text:
            raise SystemExit("%s: anchor %r not found in its row's content" % (ticker, anchor))
        out.append({"text": anchor, "href": href})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default=DEFAULT_XLSX)
    ap.add_argument("--only", help="comma-separated tickers to build (default: all)")
    a = ap.parse_args()
    only = {t.strip().upper() for t in a.only.split(",")} if a.only else None

    wb = openpyxl.load_workbook(a.xlsx)
    ids = site_ids()
    anchors = json.load(open(ANCHORS, encoding="utf-8"))

    expected = {}
    for r in list(wb["Instrument Index"].iter_rows(min_row=2, values_only=True)):
        if clean(r[1]):
            expected[clean(r[1])] = int(r[4])

    pages, order = {}, []
    for n, r in enumerate(wb["Instrument Page Content PL"].iter_rows(min_row=2, values_only=True), start=2):
        ticker, url, ctype, text = clean(r[1]), clean(r[2]), clean(r[4]), clean(r[5])
        if not ticker:
            continue
        if only and ticker.upper() not in only:
            continue
        if ctype not in TYPES:
            raise SystemExit("row %d (%s): unknown Content Type %r" % (n, ticker, ctype))
        if not text:
            raise SystemExit("row %d (%s): empty Content (PL)" % (n, ticker))
        pid = page_id(url, ids)
        if not pid:
            raise SystemExit("row %d (%s): Target Page URL %s matches no CATALOGUE symbol" % (n, ticker, url))
        if pid not in pages:
            pages[pid] = {"ticker": ticker, "source": url, "blocks": []}
            order.append(pid)
        block = {"t": TYPES[ctype], "text": text}
        links = parse_links(r[6], ticker, text, anchors)
        if links:
            block["links"] = links
        pages[pid]["blocks"].append(block)

    for pid in order:
        p = pages[pid]
        want = expected.get(p["ticker"])
        if want != len(p["blocks"]):
            raise SystemExit("%s: %d blocks built, Instrument Index expects %s" % (p["ticker"], len(p["blocks"]), want))

    body = json.dumps({k: pages[k] for k in order}, ensure_ascii=False, indent=2)
    header = (
        "/* ============================================================\n"
        "   AUTO-GENERATED — do not edit by hand.\n"
        "   Source: Instrument_Pages_PL_Content_EURUSD_to_EWJ.xlsx\n"
        "           (\"Instrument Page Content PL\", Column F verbatim)\n"
        "   Regenerate: python3 tools/build-pl-additional.py [workbook.xlsx] [--only EURUSD,…]\n"
        "   ============================================================ */\n"
        "window.PL_ADDITIONAL =\n"
    )
    with open(os.path.abspath(OUT), "w", encoding="utf-8") as f:
        f.write(header + body + ";\n")

    print("wrote %s — %d page(s)" % (os.path.abspath(OUT), len(order)))
    for pid in order:
        p = pages[pid]
        nl = sum(len(b.get("links", [])) for b in p["blocks"])
        print("  %-8s -> %-8s %d blocks (index %d) · %d link(s)" % (p["ticker"], pid, len(p["blocks"]), expected[p["ticker"]], nl))


if __name__ == "__main__":
    main()
